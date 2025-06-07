import os
from openai import OpenAI
from dotenv import load_dotenv
load_dotenv(dotenv_path="../../../.env")
from docx import Document

api_key = os.getenv("API_KEY_GROQ")

client = OpenAI(
    base_url="https://api.groq.com/openai/v1",
    api_key=api_key,
)

def read_docx_text(file_path):
    doc = Document(file_path)
    text = []
    for para in doc.paragraphs:
        text.append(para.text)
    return '\n'.join(text)

file_path = './harrypotter-pages.docx'
text = read_docx_text(file_path)

def chunk_by_paragraph(text, max_chars=3000):
    paras = text.split('\n')
    chunk = ""
    for para in paras:
        if len(chunk) + len(para) < max_chars:
            chunk += para + '\n'
        else:
            yield chunk.strip()
            chunk = para + '\n'
    if chunk.strip():
        yield chunk.strip()

translated_text = ""

for chunk in chunk_by_paragraph(text):
    prompt = (
        'Below are some English paragraphs.'
        'Please translate them into Vietnamese with a mystical tone.'
        f'"""{chunk}"""'
    )

    stream = client.chat.completions.create(
        messages=[
            {                
                "role": "system",
                "content": "You are a tool to translate from English to Vietnamese."
            },
            {
                "role": "user",
                "content": prompt,
            }
        ],
        model="llama3-8b-8192",
        temperature=0.5, 
        top_p=0.9
    )
    if stream.choices and stream.choices[0].message:
        translated_text += stream.choices[0].message.content + "\n"

print("The translated text is: ",translated_text)

def save_text_to_docx(text, output_path):
    doc = Document()
    for paragraph in text.split('\n'):
        doc.add_paragraph(paragraph)
    doc.save(output_path)

# Ví dụ
save_text_to_docx(translated_text, "harrypotter-translated.docx")
print("The translated text has been saved to harrypotter-translated.docx")